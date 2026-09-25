from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc_path = r"D:\Outputs\PaperFix\Combined_ZeroADC_CDM_Paper.docx"
doc = Document(doc_path)

# Find and update the "Data and Code Availability" section
for i, para in enumerate(doc.paragraphs):
    if "Data and Code Availability" in para.text:
        print(f"Found section at paragraph {i}: {para.text[:80]}")
        # Update the paragraph text
        para.clear()
        run = para.add_run("Data and Code Availability")
        run.bold = True
        run.font.size = Pt(11)
        
        # Add the new content
        new_text = (
            "\nThe simulation engines, test batteries, bare-metal firmware source, Renode scripts, "
            "golden vectors, and raw JSON results are publicly available at:\n"
            "• AgentForge — https://github.com/MdSadman20040812/AgentForge "
            "(falsification-first agentic automation framework with hash-based change detection)\n"
            "• ZeroADC — https://github.com/MdSadman20040812/ZeroADC "
            "(Zero-ADC analog front-end with CDM fusion, validated across simulation, Renode M33, and Arduino UNO)\n"
            "Sensor data: UCI Machine Learning Repository (#235, #374, #357)."
        )
        run2 = para.add_run(new_text)
        run2.font.size = Pt(11)
        print(f"Updated paragraph {i}")
        break
else:
    print("Section not found, appending at end")
    # Append new section at the end
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("Data and Code Availability")
    run.bold = True
    run.font.size = Pt(11)
    new_text = (
        "\nThe simulation engines, test batteries, bare-metal firmware source, Renode scripts, "
        "golden vectors, and raw JSON results are publicly available at:\n"
        "• AgentForge — https://github.com/MdSadman20040812/AgentForge "
        "(falsification-first agentic automation framework)\n"
        "• ZeroADC — https://github.com/MdSadman20040812/ZeroADC "
        "(Zero-ADC + CDM fusion, multi-platform validated)\n"
        "Sensor data: UCI Machine Learning Repository (#235, #374, #357)."
    )
    run2 = p.add_run(new_text)
    run2.font.size = Pt(11)

# Also add references to the repos in the References section
# Find the last reference or the References heading
ref_idx = None
for i, para in enumerate(doc.paragraphs):
    if para.text.strip() == "References":
        ref_idx = i
        print(f"Found References at paragraph {i}")
        break

if ref_idx:
    # Insert new references after the last existing reference
    # Find the last non-empty paragraph after References
    last_ref_idx = ref_idx + 1
    for i in range(ref_idx + 1, len(doc.paragraphs)):
        if doc.paragraphs[i].text.strip():
            last_ref_idx = i
    
    # Insert new references after the last one
    # We'll add them by inserting new paragraphs
    new_refs = [
        "[11] M. S. B. Masud, \"AgentForge: Falsification-First Agentic Automation Framework,\" GitHub Repository, 2026. [Online]. Available: https://github.com/MdSadman20040812/AgentForge",
        "[12] M. S. B. Masud, \"ZeroADC: Zero-ADC Analog Front-End with CDM Fusion for Batteryless Event-Driven Nodes,\" GitHub Repository, 2026. [Online]. Available: https://github.com/MdSadman20040812/ZeroADC"
    ]
    
    # Insert after last_ref_idx
    for ref_text in new_refs:
        new_para = doc.paragraphs[last_ref_idx]._element
        new_p = doc.add_paragraph(ref_text)
        new_para.getparent().insert(new_para.getparent().index(new_para) + 1, new_p._element)
        last_ref_idx += 1
    
    print(f"Added {len(new_refs)} new references")

doc.save(doc_path)
print(f"Saved updated document to {doc_path}")

# Verify
doc2 = Document(doc_path)
for para in doc2.paragraphs:
    if "Data and Code Availability" in para.text or "github.com" in para.text:
        print(f"VERIFIED: {para.text[:100]}")
